# Writes validation/reference/reports.json: CytoWeave's report documents
# (validation/cache/reports/, written by reference/write_reports.mjs) read back by the formats' own
# Python readers: openpyxl (the Excel workbook's cells), python-pptx (the deck's slides, tables,
# pictures and the report record part) and pypdf (the PDF's pages, text and attachments).
# validation/run.mjs (suite reports) checks what they read against CytoWeave's values. Python is
# needed only to regenerate the file, not to run the validation.
#
#   node validation/reference/write_reports.mjs
#   uv run --python 3.12 --with openpyxl==3.1.5 --with python-pptx==1.0.2 --with pypdf==5.4.0 python validation/reference/read_reports.py

import hashlib
import json
import math
from importlib.metadata import version
from pathlib import Path

import openpyxl
import pptx
import pypdf

here = Path(__file__).resolve().parent
folder = here.parent / "cache" / "reports"
manifest = json.loads((folder / "manifest.json").read_text())


def checked(name):
    data = (folder / name).read_bytes()
    entry = manifest[name]
    assert hashlib.sha256(data).hexdigest() == entry["sha256"], f"{name} is not the file the manifest lists"
    return entry["fingerprint"]


def plain(value):
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


out = {
    "about": "CytoWeave's report documents read back by openpyxl, python-pptx and pypdf; written by reference/read_reports.py.",
    "readers": {"openpyxl": version("openpyxl"), "python-pptx": version("python-pptx"), "pypdf": version("pypdf")},
    "files": {},
}

# Excel: every sheet's cells as openpyxl reads them.
book = openpyxl.load_workbook(folder / "tables.xlsx")
out["files"]["tables.xlsx"] = {
    "fingerprint": checked("tables.xlsx"),
    "sheets": [{"name": ws.title, "rows": [[plain(v) for v in row] for row in ws.iter_rows(values_only=True)]} for ws in book.worksheets],
}

# PowerPoint: each slide's tables (cell text), pictures and text boxes; the report record part.
deck = pptx.Presentation(folder / "report-sample.pptx")
slides = []
for slide in deck.slides:
    tables, pictures, texts = [], [], []
    for shape in slide.shapes:
        if shape.has_table:
            tables.append([[cell.text for cell in row.cells] for row in shape.table.rows])
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            pictures.append({"name": shape.name, "bytes": len(shape.image.blob), "type": shape.image.content_type})
        elif shape.has_text_frame:
            texts.append(shape.text_frame.text)
    slides.append({"tables": tables, "pictures": pictures, "texts": texts})
record = None
for part in deck.part.package.iter_parts():
    if str(part.partname) == "/cytoweave/report.json":
        record = json.loads(part.blob)
out["files"]["report-sample.pptx"] = {
    "fingerprint": checked("report-sample.pptx"),
    "size": [deck.slide_width, deck.slide_height],
    "slides": slides,
    "record": {"pages": len(record["pages"]), "trace": len(record["trace"])} if record else None,
}

# PDF: pages, their text as pypdf extracts it, and the attached record.
reader = pypdf.PdfReader(folder / "report-subject.pdf")
attachments = reader.attachments
record = json.loads(attachments["cytoweave-report.json"][0]) if "cytoweave-report.json" in attachments else None
out["files"]["report-subject.pdf"] = {
    "fingerprint": checked("report-subject.pdf"),
    "pages": [{"size": [float(p.mediabox.width), float(p.mediabox.height)], "text": p.extract_text()} for p in reader.pages],
    "attachments": sorted(attachments.keys()),
    "record": {"pages": len(record["pages"]), "trace": record["trace"]} if record else None,
}

(here / "reports.json").write_text(json.dumps(out, indent=1) + "\n")
print(f"wrote validation/reference/reports.json: {len(out['files'])} documents")
